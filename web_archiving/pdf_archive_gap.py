#!/usr/bin/env python3
"""
pdf_archive_gap.py — Find PDFs on uni.edu that are not yet captured in UNI's Archive-It collections.

Produces these CSVs (in --outdir):
  uni_pdfs.csv         every PDF link found on the live uni.edu domain (incl. subdomains), cumulative across runs
  archiveit_pdfs.csv   every PDF captured in your Archive-It collection(s)
  pdf_comparison.xlsx  the review workbook — one row per PDF with a status:
                         NEEDS_ARCHIVING      live on uni.edu, not in Archive-It
                         ARCHIVED             live and already captured
                         ARCHIVED_NOT_LIVE    captured, but not found on the live site
                       plus three yellow columns for a human reviewer (with dropdowns):
                         appraisal              add | ignore   (blank = not yet reviewed)
                         target_collection_id   Archive-It collection the PDF should go into
                         seed_type              One Page (seed = the PDF; blank means this) | One Page+ (seed = the
                                                page it was found on, captured with its links)
                       and seeded_on, which the script fills in when a row is exported to a seed list.
                       Reviewer entries are read back and preserved on every rerun, so keep
                       editing the same file. Sheets: Review, Collections, Instructions.
  seeds/<timestamp>/   seed lists built from rows marked "add", one per collection and seed type:
                       collection_<ID>_one-page_seeds.txt / collection_<ID>_one-page-plus_seeds.txt
                       (one URL per line, ≤1,000 per file to fit Archive-It's Add Seeds dialog),
                       seed_summary.csv and README.txt with the Archive-It steps.

Crawl state (also in --outdir) lets you rerun without re-crawling pages:
  crawled_pages.csv    log of every page fetched: key, url, crawled_at, status
  crawl_queue.txt      pages discovered but not yet crawled (where the next run picks up)

Requires Python 3.8+ and openpyxl (one-time install: python -m pip install openpyxl).

Typical cycle:
  1. python pdf_archive_gap.py               crawl + compare, update pdf_comparison.xlsx
  2. reviewers fill in appraisal / target_collection_id / seed_type, save and close the workbook
  3. python pdf_archive_gap.py --seeds-only  write seed lists for rows marked add (no crawling)
  4. paste each seeds file into Archive-It (Seeds tab → Add Seeds, seed type as named in the file) and crawl
  A full run (step 1) also exports any reviewed-but-unexported rows, so step 3 is optional.

Usage:
  python pdf_archive_gap.py                          # first run, or resume/continue from saved state
  python pdf_archive_gap.py --max-pages 2000         # crawl in batches: run it again to do the next 2000
  python pdf_archive_gap.py --recrawl-days 90        # also re-crawl pages last crawled 90+ days ago
  python pdf_archive_gap.py --fresh                  # ignore saved state and start over
  python pdf_archive_gap.py --collections 1234 5678  # name Archive-It collections instead of auto-discovering
  python pdf_archive_gap.py --default-collection 1234  # pre-fill target_collection_id on unreviewed gap rows
  python pdf_archive_gap.py --seeds-only             # export seed lists from the reviewed workbook only
  python pdf_archive_gap.py --seeds-only --reseed    # re-export everything marked add, even if seeded before

Rerun behaviour:
  * Pages already in crawled_pages.csv are skipped (unless older than --recrawl-days).
  * The saved queue is resumed first, then the home page and sitemaps are re-checked for new pages.
  * PDFs found in earlier runs are kept; new ones are added with today's date in first_seen.
  * Pages that failed to load are NOT marked as crawled, so they're retried next run.
  * Ctrl-C saves state before exiting, so an interrupted run can be resumed.
  * Archive-It is re-queried every run (it's quick and its contents change).

Notes:
  * PDFs are detected by ".pdf" in the URL path. PDFs served without that extension won't be seen.
  * A page skipped as already crawled isn't re-read, so a NEW link added to an OLD page is only
    found when that page is re-crawled — use --recrawl-days for periodic full refreshes.
  * Archive-It lookup uses the public Wayback CDX endpoint for each collection, so only
    public collections are visible. Private collections need a WASAPI/partner export instead.
"""
import argparse, csv, gzip, json, os, re, sys, time, urllib.parse, urllib.request, xml.etree.ElementTree as ET
from collections import deque
from datetime import datetime, timedelta
from html.parser import HTMLParser

UA = "UNI-PDF-Archive-Audit/1.0 (+contact: web archiving staff)"
ROOT_DOMAIN = "uni.edu"
SKIP_EXT = re.compile(r"\.(jpe?g|png|gif|svg|webp|ico|css|js|mp4|mp3|mov|zip|docx?|xlsx?|pptx?|woff2?|ttf)$", re.I)
TS_FMT = "%Y-%m-%d %H:%M:%S"


# ---------------------------------------------------------------- helpers
def fetch(url, timeout=30, retries=2):
    for attempt in range(retries + 1):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA})
            with urllib.request.urlopen(req, timeout=timeout) as r:
                data = r.read()
                if r.headers.get("Content-Encoding") == "gzip" or url.endswith(".gz"):
                    try: data = gzip.decompress(data)
                    except OSError: pass
                return r.geturl(), r.headers.get("Content-Type", ""), data
        except urllib.error.HTTPError as e:
            if e.code == 429 and attempt < retries:
                time.sleep(30 * (attempt + 1)); continue
            return None, None, None
        except Exception:
            if attempt < retries: time.sleep(3); continue
            return None, None, None


def on_domain(url):
    host = (urllib.parse.urlsplit(url).hostname or "").lower()
    return host == ROOT_DOMAIN or host.endswith("." + ROOT_DOMAIN)


def is_pdf(url):
    return urllib.parse.urlsplit(url).path.lower().endswith(".pdf")


def normalize(url):
    """Comparison key: ignore scheme, leading www., default ports, fragments, trailing '/', case of host."""
    p = urllib.parse.urlsplit(url.strip())
    host = (p.hostname or "").lower()
    if host.startswith("www."): host = host[4:]
    path = urllib.parse.unquote(p.path or "/")
    path = urllib.parse.quote(path, safe="/-._~%!$&'()*+,;=:@")
    if len(path) > 1: path = path.rstrip("/")
    return f"{host}{path}" + (f"?{p.query}" if p.query else "")


class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__(); self.links = []
    def handle_starttag(self, tag, attrs):
        if tag in ("a", "link", "iframe", "embed", "object", "area"):
            for k, v in attrs:
                if k in ("href", "src", "data") and v: self.links.append(v)


# ---------------------------------------------------------------- crawl state
def state_paths(d):
    return {"log": os.path.join(d, "crawled_pages.csv"),
            "queue": os.path.join(d, "crawl_queue.txt"),
            "pdfs": os.path.join(d, "uni_pdfs.csv")}


def load_state(d):
    """Return (seen: key -> last successful crawl time, pdfs: url -> info, queue: [urls])."""
    p = state_paths(d)
    seen, pdfs, queue = {}, {}, []
    if os.path.exists(p["log"]):
        with open(p["log"], newline="") as f:
            for r in csv.DictReader(f):
                if r.get("status") == "ok" and r["crawled_at"] > seen.get(r["key"], ""):
                    seen[r["key"]] = r["crawled_at"]
    if os.path.exists(p["pdfs"]):
        with open(p["pdfs"], newline="") as f:
            for r in csv.DictReader(f):
                pdfs[r["pdf_url"]] = {"found_on": r.get("found_on_page", ""), "first_seen": r.get("first_seen", "")}
    if os.path.exists(p["queue"]):
        with open(p["queue"]) as f:
            queue = [l.strip() for l in f if l.strip()]
    return seen, pdfs, queue


def save_state(d, pdfs, queue, done):
    p = state_paths(d)
    pending, keys = [], set()
    for u in queue:                      # dedupe and drop anything finished
        k = normalize(u)
        if k not in keys and not done(k):
            keys.add(k); pending.append(u)
    tmp = p["queue"] + ".tmp"
    with open(tmp, "w") as f: f.write("\n".join(pending) + ("\n" if pending else ""))
    os.replace(tmp, p["queue"])
    tmp = p["pdfs"] + ".tmp"
    with open(tmp, "w", newline="") as f:
        w = csv.writer(f); w.writerow(["pdf_url", "found_on_page", "first_seen"])
        for u, info in sorted(pdfs.items()): w.writerow([u, info["found_on"], info["first_seen"]])
    os.replace(tmp, p["pdfs"])
    return len(pending)


# ---------------------------------------------------------------- 1. live site crawl
def sitemap_urls():
    urls, queue, seen = set(), deque([f"https://{ROOT_DOMAIN}/sitemap.xml"]), set()
    _, _, robots = fetch(f"https://{ROOT_DOMAIN}/robots.txt")
    if robots:
        for line in robots.decode("utf-8", "ignore").splitlines():
            if line.lower().startswith("sitemap:"): queue.append(line.split(":", 1)[1].strip())
    while queue:
        sm = queue.popleft()
        if sm in seen: continue
        seen.add(sm)
        _, _, data = fetch(sm)
        if not data: continue
        try: root = ET.fromstring(data)
        except ET.ParseError: continue
        for loc in root.iter():
            if loc.tag.endswith("loc") and loc.text:
                u = loc.text.strip()
                if root.tag.endswith("sitemapindex"): queue.append(u)
                else: urls.add(u)
    return urls


def crawl(outdir, max_pages, delay, recrawl_days, fresh):
    now = datetime.now()
    run_ts = now.strftime(TS_FMT)
    cutoff = (now - timedelta(days=recrawl_days)).strftime(TS_FMT) if recrawl_days is not None else None
    p = state_paths(outdir)

    if fresh:
        for f in p.values():
            if os.path.exists(f): os.remove(f)
    seen, pdfs, saved_queue = load_state(outdir)
    prior_pages, prior_pdfs = len(seen), len(pdfs)
    if prior_pages:
        print(f"  resuming: {prior_pages} pages already crawled, {prior_pdfs} PDFs known, "
              f"{len(saved_queue)} pages queued from last run", file=sys.stderr)

    def done(key):
        ts = seen.get(key)
        return bool(ts) and (cutoff is None or ts > cutoff)

    # saved queue first (continue where we left off), then home page + sitemaps to pick up new pages
    seeds = {f"https://{ROOT_DOMAIN}/"} | sitemap_urls()
    queue = deque(saved_queue)
    queue.extend(sorted(seeds))
    for s in seeds:
        if is_pdf(s) and s not in pdfs: pdfs[s] = {"found_on": "sitemap", "first_seen": run_ts}

    new_log = not os.path.exists(p["log"])
    log = open(p["log"], "a", newline="")
    lw = csv.writer(log)
    if new_log: lw.writerow(["key", "url", "crawled_at", "status"])

    this_run, failed, pages, skipped, interrupted = set(), [], 0, 0, False
    try:
        while queue and pages < max_pages:
            url = queue.popleft()
            if is_pdf(url) or SKIP_EXT.search(urllib.parse.urlsplit(url).path): continue
            key = normalize(url)
            if key in this_run: continue
            if done(key): skipped += 1; continue
            this_run.add(key)

            final, ctype, body = fetch(url)
            pages += 1
            ok = body is not None
            lw.writerow([key, url, datetime.now().strftime(TS_FMT), "ok" if ok else "error"])
            if ok: seen[key] = datetime.now().strftime(TS_FMT)
            else: failed.append(url)          # keep it queued so the next run retries it

            if ok and "html" in (ctype or ""):
                lp = LinkParser()
                try: lp.feed(body.decode("utf-8", "ignore"))
                except Exception: lp.links = []
                for href in lp.links:
                    absu = urllib.parse.urljoin(final or url, href).split("#")[0]
                    if not absu.startswith("http") or not on_domain(absu): continue
                    if is_pdf(absu):
                        if absu not in pdfs: pdfs[absu] = {"found_on": final or url, "first_seen": run_ts}
                    else:
                        k = normalize(absu)
                        if k not in this_run and not done(k): queue.append(absu)

            if pages % 100 == 0:
                log.flush(); save_state(outdir, pdfs, list(queue) + failed, done)
                print(f"  crawled {pages} pages this run, {len(pdfs)} PDFs, {len(queue)} queued "
                      f"(checkpoint saved)", file=sys.stderr)
            time.sleep(delay)
    except KeyboardInterrupt:
        interrupted = True
        print("\n  interrupted — saving crawl state…", file=sys.stderr)
    finally:
        log.close()
        remaining = save_state(outdir, pdfs, list(queue) + failed, done)

    print(f"Crawl: {pages} pages fetched this run, {skipped} skipped as already crawled, "
          f"{len(pdfs) - prior_pdfs} new PDFs ({len(pdfs)} total)", file=sys.stderr)
    if remaining:
        print(f"  {remaining} pages still queued — rerun to continue.", file=sys.stderr)
    if interrupted: sys.exit(1)
    return pdfs, run_ts


# ---------------------------------------------------------------- 2. Archive-It
def discover_collections(org="uni"):
    ids = set()
    for page in range(1, 20):
        _, _, body = fetch(f"https://archive-it.org/home/{org}?page={page}")
        if not body: break
        found = set(re.findall(r"/collections/(\d+)", body.decode("utf-8", "ignore")))
        if not found - ids: break
        ids |= found
        time.sleep(2)
    return sorted(ids, key=int)


def archiveit_pdfs(collections):
    out = {}  # normalized -> row
    for coll in collections:
        for match in ("domain", "prefix"):
            q = urllib.parse.urlencode({"url": ROOT_DOMAIN if match == "domain" else f"{ROOT_DOMAIN}/",
                                        "matchType": match, "filter": "mimetype:application/pdf",
                                        "output": "json", "fl": "original,timestamp,statuscode"})
            _, _, body = fetch(f"https://wayback.archive-it.org/{coll}/timemap/cdx?{q}", timeout=300)
            if not body: continue
            text = body.decode("utf-8", "ignore").strip()
            if text.startswith("["):
                data = json.loads(text); rows = data[1:] if data and data[0][0] == "original" else data
            else:
                rows = [l.split() for l in text.splitlines() if l.strip()]
            for r in rows:
                if not r: continue
                orig = r[0] if r[0].startswith("http") else next((x for x in r if x.startswith("http")), None)
                if not orig: continue
                ts = next((x for x in r if re.fullmatch(r"\d{14}", x)), "")
                k = normalize(orig)
                rec = out.setdefault(k, {"url": orig, "collections": set(), "first": ts, "last": ts, "captures": 0})
                rec["collections"].add(coll); rec["captures"] += 1
                if ts and (not rec["first"] or ts < rec["first"]): rec["first"] = ts
                if ts > rec["last"]: rec["last"] = ts
            print(f"  collection {coll}: {len(rows)} PDF captures ({match} match)", file=sys.stderr)
            if rows: break
        time.sleep(2)
    return out


# ---------------------------------------------------------------- 3. compare → review workbook → seed lists
WORKBOOK = "pdf_comparison.xlsx"
LEGACY_CSV = "pdf_comparison.csv"
SHEET = "Review"
COLUMNS = ["status", "appraisal", "target_collection_id", "seed_type", "pdf_url", "found_on_page", "first_seen",
           "new_this_run", "archiveit_collections", "last_capture", "seeded_on"]
INPUT_COLS = ("appraisal", "target_collection_id", "seed_type")    # the yellow, reviewer-edited columns
APPRAISAL_VALUES = {"add", "ignore"}
SEED_TYPES = ("One Page", "One Page+")
DEFAULT_SEED_TYPE = "One Page"     # used when seed_type is left blank
STATUS_ORDER = {"NEEDS_ARCHIVING": 0, "ARCHIVED_NOT_LIVE": 1, "ARCHIVED": 2}
SEED_BATCH = 1000          # Archive-It's "Add Seeds" dialog accepts up to 1,000 lines at a time
MAX_HYPERLINKS = 65000     # Excel's per-sheet hyperlink limit is 65,530


def fmt_ts(ts):
    return f"{ts[:4]}-{ts[4:6]}-{ts[6:8]}" if ts else ""


def need_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        sys.exit("The Excel review file needs openpyxl. Install it once with:\n"
                 "    python -m pip install openpyxl")


def as_text(v):
    """Cell value -> clean string (handles Excel turning IDs into floats and dates into datetimes)."""
    if v is None: return ""
    if hasattr(v, "strftime"): return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer(): v = int(v)
    s = str(v).strip()
    return s[:-2] if s.endswith(".0") and s[:-2].isdigit() else s


def norm_seed_type(v):
    """'one page', 'OnePage', 'One Page +', 'one page plus links' ... -> 'One Page' / 'One Page+'.
    Blank -> '' (treated as DEFAULT_SEED_TYPE). Unrecognised -> None."""
    s = re.sub(r"[\s_-]", "", str(v or "")).lower()
    if not s: return ""
    if s == "onepage": return "One Page"
    if s in ("onepage+", "onepageplus", "onepage+links", "onepagepluslinks"): return "One Page+"
    return None


def ensure_writable(path):
    """Excel locks files it has open; catch that before hours of crawling, not after."""
    if not os.path.exists(path): return
    try:
        with open(path, "r+b"): pass
    except PermissionError:
        sys.exit(f"{path} is open in another program (probably Excel). Save and close it, then rerun.")


def read_workbook_rows(path):
    """Rows of the Review sheet as dicts keyed by header; '_row' holds the sheet row number."""
    openpyxl = need_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [as_text(h) for h in next(it, [])]
    rows = []
    for n, vals in enumerate(it, start=2):
        r = {h: as_text(v) for h, v in zip(header, vals) if h}
        if r.get("pdf_url"):
            r["_row"] = n
            rows.append(r)
    wb.close()
    return rows


def load_reviews(outdir):
    """Reviewer entries (appraisal, target_collection_id) + script-managed seeded_on from the last run.
    Reads pdf_comparison.xlsx, or the older pdf_comparison.csv if no workbook exists yet."""
    xlsx, legacy = os.path.join(outdir, WORKBOOK), os.path.join(outdir, LEGACY_CSV)
    if os.path.exists(xlsx):
        src, rows = WORKBOOK, read_workbook_rows(xlsx)
    elif os.path.exists(legacy):
        with open(legacy, newline="", encoding="utf-8-sig") as f:   # utf-8-sig: tolerate Excel's BOM
            src, rows = LEGACY_CSV, [{k: as_text(v) for k, v in r.items() if k} for r in csv.DictReader(f)]
    else:
        return {}, [], None
    reviews, bad = {}, []
    for r in rows:
        url = r.get("pdf_url", "")
        if not url: continue
        appraisal = r.get("appraisal", "").lower()
        target = r.get("target_collection_id", "")
        stype_raw = r.get("seed_type", "")
        stype = norm_seed_type(stype_raw)
        seeded = r.get("seeded_on", "")
        if appraisal and appraisal not in APPRAISAL_VALUES: bad.append((url, f"appraisal '{appraisal}'"))
        if stype is None: bad.append((url, f"seed_type '{stype_raw}'")); stype = stype_raw   # keep so reviewer sees it
        if appraisal or target or stype or seeded:
            reviews[normalize(url)] = {"appraisal": appraisal, "target_collection_id": target,
                                       "seed_type": stype, "seeded_on": seeded}
    return reviews, bad, src


def build_rows(pdfs, arch, reviews, run_ts, default_collection):
    live = {normalize(u): (u, info) for u, info in pdfs.items()}
    rows = []
    for k in set(live) | set(arch):
        lv, ar = live.get(k), arch.get(k)
        status = "ARCHIVED" if lv and ar else "NEEDS_ARCHIVING" if lv else "ARCHIVED_NOT_LIVE"
        rv = reviews.get(k, {})
        target = rv.get("target_collection_id", "")
        if not target and status == "NEEDS_ARCHIVING" and default_collection:
            target = str(default_collection)
        rows.append({
            "status": status,
            "appraisal": rv.get("appraisal", ""),
            "target_collection_id": target,
            "seed_type": rv.get("seed_type", ""),
            "pdf_url": lv[0] if lv else ar["url"],
            "found_on_page": lv[1]["found_on"] if lv else "",
            "first_seen": lv[1]["first_seen"][:10] if lv else "",
            "new_this_run": "yes" if lv and lv[1]["first_seen"] == run_ts else "",
            "archiveit_collections": ";".join(sorted(ar["collections"])) if ar else "",
            "last_capture": fmt_ts(ar["last"]) if ar else "",
            "seeded_on": rv.get("seeded_on", ""),
        })
    # within each status: unreviewed first, then newly found, then alphabetical
    rows.sort(key=lambda r: (STATUS_ORDER[r["status"]], r["appraisal"] in APPRAISAL_VALUES,
                             r["new_this_run"] != "yes", r["pdf_url"]))
    return rows


def write_workbook(path, rows, collections):
    openpyxl = need_openpyxl()
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.utils import get_column_letter

    F = Font(name="Arial", size=10)
    FB = Font(name="Arial", size=10, bold=True)
    FLINK = Font(name="Arial", size=10, color="0563C1", underline="single")
    FAUTO = Font(name="Arial", size=10, italic=True, color="808080")
    HEAD = PatternFill("solid", fgColor="D9D9D9")
    INPUT_HEAD = PatternFill("solid", fgColor="FFD966")
    INPUT = PatternFill("solid", fgColor="FFF2CC")
    col = {c: i + 1 for i, c in enumerate(COLUMNS)}
    L = {c: get_column_letter(i) for c, i in col.items()}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(COLUMNS)
    links = 0
    for r in rows:
        vals = []
        for c in COLUMNS:
            v = r.get(c, "")
            if c == "target_collection_id" and str(v).isdigit(): v = int(v)   # numeric, to match the dropdown
            vals.append(v)
        ws.append(vals)
        i = ws.max_row
        for c in COLUMNS:
            ws.cell(i, col[c]).font = FAUTO if c == "seeded_on" else F
        for c in INPUT_COLS:
            ws.cell(i, col[c]).fill = INPUT
        ws.cell(i, col["target_collection_id"]).alignment = Alignment(horizontal="left")
        if links < MAX_HYPERLINKS:
            cell = ws.cell(i, col["pdf_url"])
            cell.hyperlink, cell.font = r["pdf_url"], FLINK
            links += 1
    for c in COLUMNS:
        h = ws.cell(1, col[c])
        h.font, h.fill = FB, (INPUT_HEAD if c in INPUT_COLS else HEAD)
        h.alignment = Alignment(vertical="center")
    widths = {"status": 20, "appraisal": 11, "target_collection_id": 20, "seed_type": 12, "pdf_url": 70,
              "found_on_page": 45, "first_seen": 12, "new_this_run": 13, "archiveit_collections": 21,
              "last_capture": 13, "seeded_on": 12}
    for c, w in widths.items(): ws.column_dimensions[L[c]].width = w
    ws.freeze_panes = "E2"
    last = max(ws.max_row, 2)
    ws.auto_filter.ref = f"A1:{L['seeded_on']}{last}"

    # dropdowns
    dv = DataValidation(type="list", formula1='"add,ignore"', allow_blank=True, showDropDown=False,
                        showErrorMessage=True, errorTitle="appraisal", error="Choose add or ignore, or leave blank.")
    dv.add(f"{L['appraisal']}2:{L['appraisal']}{last}")
    ws.add_data_validation(dv)
    dv3 = DataValidation(type="list", formula1='"' + ",".join(SEED_TYPES) + '"', allow_blank=True,
                         showErrorMessage=True, errorTitle="seed_type",
                         error=f"Choose {' or '.join(SEED_TYPES)}, or leave blank for {DEFAULT_SEED_TYPE}.")
    dv3.add(f"{L['seed_type']}2:{L['seed_type']}{last}")
    ws.add_data_validation(dv3)

    cs = wb.create_sheet("Collections")
    cs.append(["collection_id", "source"])
    for cid, src in collections:
        cs.append([int(cid) if str(cid).isdigit() else cid, src])
    for row in cs.iter_rows():
        for c in row: c.font = F
    for c in cs[1]: c.font, c.fill = FB, HEAD
    cs.column_dimensions["A"].width, cs.column_dimensions["B"].width = 16, 40
    if collections:
        dv2 = DataValidation(type="list", formula1=f"=Collections!$A$2:$A${len(collections) + 1}", allow_blank=True,
                             showErrorMessage=True, errorStyle="warning", errorTitle="target_collection_id",
                             error="This ID isn't on the Collections sheet. Click Yes to use it anyway.")
        dv2.add(f"{L['target_collection_id']}2:{L['target_collection_id']}{last}")
        ws.add_data_validation(dv2)

    # highlight decisions
    rng = f"A2:{L['seeded_on']}{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${L["appraisal"]}2="add"'],
                                                   fill=PatternFill("solid", fgColor="E2EFDA")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${L["appraisal"]}2="ignore"'],
                                                   font=Font(color="A6A6A6")))

    add_instructions(wb.create_sheet("Instructions"), F, FB, HEAD, INPUT)

    tmp = path + ".tmp"
    wb.save(tmp)
    os.replace(tmp, path)


def add_instructions(ws, F, FB, HEAD, INPUT):
    lines = [
        ("How to review", "title"),
        ("Work on the Review sheet. Only edit the three yellow columns; everything else is regenerated on each run.", ""),
        ("", ""),
        ("appraisal", "head"),
        ("add — this PDF should be captured in Archive-It", ""),
        ("ignore — don't capture it", ""),
        ("(blank) — not reviewed yet. Unreviewed rows sort to the top of each status group.", ""),
        ("", ""),
        ("target_collection_id", "head"),
        ("The Archive-It collection the PDF should go into. The dropdown lists known collections (Collections sheet);", ""),
        ("you can type a different ID and click Yes at the warning.", ""),
        ("", ""),
        ("seed_type", "head"),
        ("One Page — the seed is the PDF itself. Archive-It captures just that file. (Blank means One Page.)", ""),
        ("One Page+ — the seed is the page the PDF was found on (found_on_page). Archive-It captures that page plus", ""),
        ("   everything it links to one hop out, so the PDF is archived in context, along with the page's other links.", ""),
        ("   Several PDFs found on the same page become one seed.", ""),
        ("", ""),
        ("seeded_on (grey, filled in by the script)", "head"),
        ("The date this PDF was exported to a seed list. Rows with a date aren't exported again.", ""),
        ("Clear the date to export that PDF again (or run with --reseed to re-export everything marked add).", ""),
        ("", ""),
        ("status", "head"),
        ("NEEDS_ARCHIVING — on uni.edu, not found in Archive-It", ""),
        ("ARCHIVED_NOT_LIVE — in Archive-It, no longer linked on uni.edu (can't be seeded: there's nothing live to crawl)", ""),
        ("ARCHIVED — on uni.edu and already captured (mark add only if you want a fresh capture)", ""),
        ("", ""),
        ("Example of a reviewed row", "head"),
        ("EXAMPLE", "example"),
        ("", ""),
        ("Workflow", "head"),
        ("1. Run:  python pdf_archive_gap.py   (crawls, compares, updates this workbook)", ""),
        ("2. Review the yellow columns, then save and close this file.", ""),
        ("3. Run:  python pdf_archive_gap.py --seeds-only   (writes seed lists to seeds/<date>/, one file per collection and seed type)", ""),
        ("4. In Archive-It: open the collection → Seeds tab → Add Seeds → paste the file's contents (one URL per line,", ""),
        ("   max 1,000 per paste — longer lists are split into _part files) → set Seed Type to the one in the file name", ""),
        ("   (one-page or one-page-plus) → choose visibility → Add Seeds.", ""),
        ("5. Crawl the new seeds. The next run will show them as ARCHIVED once captures are indexed.", ""),
    ]
    ws.column_dimensions["A"].width = 22
    for c in "BCDEFGHIJ": ws.column_dimensions[c].width = 18
    from openpyxl.styles import Font
    row = 1
    for text, kind in lines:
        if kind == "example":
            ex = ["NEEDS_ARCHIVING", "add", 12345, "One Page", "https://uni.edu/registrar/forms/example.pdf",
                  "https://uni.edu/registrar/forms", "2026-10-26", "yes", "", "", ""]
            for j, h in enumerate(COLUMNS, start=1):
                c = ws.cell(row, j, h); c.font, c.fill = FB, HEAD
            for j, v in enumerate(ex, start=1):
                c = ws.cell(row + 1, j, v); c.font = F
                if COLUMNS[j - 1] in INPUT_COLS: c.fill = INPUT
            row += 2
            continue
        c = ws.cell(row, 1, text if text else None)
        c.font = Font(name="Arial", size=14, bold=True) if kind == "title" else FB if kind == "head" else F
        row += 1


def export_seeds(outdir, rows, reseed=False):
    """Write seed lists for rows marked 'add': one file per collection AND seed type, because Archive-It's
    Add Seeds dialog applies a single seed type to each pasted batch. Sets seeded_on on exported rows.
      One Page   -> seed is the PDF URL
      One Page+  -> seed is the page the PDF was found on (deduplicated)
    Returns {(collection_id, seed_type): {seed_url: [(pdf_url, status), ...]}} (empty if nothing exported)."""
    today = datetime.now().strftime("%Y-%m-%d")
    groups, picked = {}, []
    skipped = {"no target_collection_id": [], "no longer on uni.edu": [],
               "unrecognised seed_type": [], "One Page+ but no linking page is known": []}
    for r in rows:
        if r.get("appraisal", "").lower() != "add": continue
        if r.get("seeded_on") and not reseed: continue
        url = r["pdf_url"]
        if r.get("status") == "ARCHIVED_NOT_LIVE": skipped["no longer on uni.edu"].append(url); continue
        coll = as_text(r.get("target_collection_id"))
        if not coll: skipped["no target_collection_id"].append(url); continue
        stype = norm_seed_type(r.get("seed_type"))
        if stype is None: skipped["unrecognised seed_type"].append(f"{url} ({r.get('seed_type')})"); continue
        stype = stype or DEFAULT_SEED_TYPE
        if stype == "One Page+":
            seed = r.get("found_on_page", "")
            if not seed.startswith("http"):        # e.g. PDF came from the sitemap, not a page
                skipped["One Page+ but no linking page is known"].append(url); continue
        else:
            seed = url
        groups.setdefault((coll, stype), {}).setdefault(seed, []).append((url, r.get("status", "")))
        picked.append(r)

    for why, urls in skipped.items():
        for u in urls:
            print(f"  seed lists: skipped ({why}): {u}", file=sys.stderr)
    if not groups:
        print("  seed lists: nothing new marked 'add' with a collection ID — no files written", file=sys.stderr)
        return {}

    now = datetime.now()
    stamp = now.strftime("%Y-%m-%d_%H%M%S")
    d = os.path.join(outdir, "seeds", stamp)
    os.makedirs(d, exist_ok=True)
    summary, readme = [], [f"Archive-It seed lists exported {now:%Y-%m-%d %H:%M}", ""]
    order = lambda k: (not k[0].isdigit(), k[0].zfill(12), SEED_TYPES.index(k[1]))
    for coll, stype in sorted(groups, key=order):
        seeds = list(groups[(coll, stype)])
        safe = re.sub(r"[^\w-]", "_", coll)
        tag = "one-page-plus" if stype == "One Page+" else "one-page"
        chunks = [seeds[i:i + SEED_BATCH] for i in range(0, len(seeds), SEED_BATCH)]
        for n, chunk in enumerate(chunks, start=1):
            name = f"collection_{safe}_{tag}_seeds" + (f"_part{n}" if len(chunks) > 1 else "") + ".txt"
            with open(os.path.join(d, name), "w", encoding="utf-8", newline="\n") as f:
                f.write("\n".join(chunk) + "\n")
            n_pdfs = sum(len(groups[(coll, stype)][s]) for s in chunk)
            extra = f", covering {n_pdfs} PDFs" if stype == "One Page+" else ""
            readme.append(f"{name}\n    collection {coll} | Seed Type: {stype} | {len(chunk)} "
                          f"seed{'s' if len(chunk) != 1 else ''}{extra}")
            for s in chunk:
                for pdf, status in groups[(coll, stype)][s]:
                    summary.append([coll, stype, s, pdf, status, name])
    readme += ["", "To add them in Archive-It, for each file:",
               "  1. Open the collection named above, go to the Seeds tab and click Add Seeds.",
               "  2. Paste the file's contents (one URL per line; each file holds at most 1,000, the dialog's limit).",
               "  3. Set Seed Type to the type named above for that file. Add Seeds applies one seed type to the",
               "     whole batch, which is why lists are split by type.",
               "       One Page  : the seeds are the PDFs themselves; just each file is captured.",
               "       One Page+ : the seeds are the pages the PDFs were found on; each page is captured with",
               "                   everything it links to one hop out, including the PDFs.",
               "  4. Choose visibility and crawl frequency, then click Add Seeds.",
               "  5. Crawl the new seeds.",
               "",
               "seed_summary.csv lists every PDF with the seed that covers it, its collection, seed type and file."]
    write_csv(os.path.join(d, "seed_summary.csv"),
              ["collection_id", "seed_type", "seed_url", "pdf_url", "status", "file"], summary)
    with open(os.path.join(d, "README.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(readme) + "\n")

    for r in picked: r["seeded_on"] = today
    total = sum(len(v) for v in groups.values())
    print(f"  seed lists: {total} seeds ({len(picked)} PDFs) in {len(groups)} list(s) written to {d}", file=sys.stderr)
    return groups


def mark_seeded_in_workbook(path, rows):
    """--seeds-only: write seeded_on dates back into the existing workbook, leaving everything else alone."""
    openpyxl = need_openpyxl()
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    header = [as_text(c.value) for c in ws[1]]
    if "seeded_on" not in header:
        ws.cell(1, len(header) + 1, "seeded_on")
        header.append("seeded_on")
    ci = header.index("seeded_on") + 1
    for r in rows:
        if r.get("seeded_on"): ws.cell(r["_row"], ci, r["seeded_on"])
    tmp = path + ".tmp"
    wb.save(tmp)
    os.replace(tmp, path)


def write_csv(path, header, rows):
    tmp = path + ".tmp"
    with open(tmp, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(header); w.writerows(rows)
    os.replace(tmp, path)


def main():
    for stream in (sys.stdout, sys.stderr):      # never crash on a console that can't show a character
        try: stream.reconfigure(errors="replace")
        except AttributeError: pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--collections", nargs="*", help="Archive-It collection IDs (default: discover from archive-it.org/home/uni)")
    ap.add_argument("--max-pages", type=int, default=10000, help="max pages to fetch THIS run (default 10000)")
    ap.add_argument("--delay", type=float, default=0.3, help="seconds between page requests")
    ap.add_argument("--recrawl-days", type=int, default=None,
                    help="re-crawl pages last crawled this many days ago or more (default: never re-crawl)")
    ap.add_argument("--fresh", action="store_true", help="delete saved crawl state and start from scratch")
    ap.add_argument("--default-collection", default=None,
                    help="pre-fill target_collection_id with this ID on NEEDS_ARCHIVING rows that have none yet")
    ap.add_argument("--seeds-only", action="store_true",
                    help="skip crawling; just export seed lists from the reviewed pdf_comparison.xlsx")
    ap.add_argument("--reseed", action="store_true", help="also re-export rows that already have a seeded_on date")
    ap.add_argument("--outdir", default=".")
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    need_openpyxl()
    wb_path = os.path.join(a.outdir, WORKBOOK)
    ensure_writable(wb_path)

    if a.seeds_only:
        if not os.path.exists(wb_path):
            sys.exit(f"No {WORKBOOK} in {a.outdir} yet — run once without --seeds-only first.")
        rows = read_workbook_rows(wb_path)
        bad = {r["appraisal"] for r in rows if r.get("appraisal") and r["appraisal"].lower() not in APPRAISAL_VALUES}
        if bad: print(f"  warning: ignoring appraisal values other than add/ignore: {', '.join(sorted(bad))}", file=sys.stderr)
        if export_seeds(a.outdir, rows, a.reseed):
            mark_seeded_in_workbook(wb_path, rows)
            print(f"  recorded seeded_on dates in {WORKBOOK}", file=sys.stderr)
        return

    print("1/4 Crawling uni.edu for PDFs…", file=sys.stderr)
    pdfs, run_ts = crawl(a.outdir, a.max_pages, a.delay, a.recrawl_days, a.fresh)

    print("2/4 Querying Archive-It…", file=sys.stderr)
    colls = a.collections or discover_collections()
    if not colls: sys.exit("No Archive-It collection IDs found — pass them with --collections")
    print(f"  collections: {', '.join(colls)}", file=sys.stderr)
    arch = archiveit_pdfs(colls)
    write_csv(os.path.join(a.outdir, "archiveit_pdfs.csv"),
              ["pdf_url", "collections", "captures", "first_capture", "last_capture"],
              [[r["url"], ";".join(sorted(r["collections"])), r["captures"], fmt_ts(r["first"]), fmt_ts(r["last"])]
               for r in sorted(arch.values(), key=lambda r: r["url"])])

    print("3/4 Comparing…", file=sys.stderr)
    reviews, bad, src = load_reviews(a.outdir)
    if reviews:
        print(f"  carried over {len(reviews)} reviewed rows from {src}", file=sys.stderr)
    for url, val in bad:
        print(f"  warning: {val} isn't one of the dropdown choices — kept as-is for {url}", file=sys.stderr)
    rows = build_rows(pdfs, arch, reviews, run_ts, a.default_collection)

    print("4/4 Exporting seed lists and writing the review workbook…", file=sys.stderr)
    export_seeds(a.outdir, rows, a.reseed)
    known = {c: "Archive-It (searched this run)" for c in colls}
    if a.default_collection: known.setdefault(str(a.default_collection), "--default-collection")
    for r in rows:
        if r["target_collection_id"]: known.setdefault(r["target_collection_id"], "entered by a reviewer")
    collections = sorted(known.items(), key=lambda kv: (not kv[0].isdigit(), kv[0].zfill(12)))
    try:
        write_workbook(wb_path, rows, collections)
    except PermissionError:
        alt = os.path.join(a.outdir, f"pdf_comparison_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
        write_workbook(alt, rows, collections)
        print(f"  {WORKBOOK} was locked by another program — wrote {alt} instead. Rename it to {WORKBOOK} "
              f"before the next run so reviews and seeded_on dates carry over.", file=sys.stderr)
    legacy = os.path.join(a.outdir, LEGACY_CSV)
    if src == LEGACY_CSV and os.path.exists(legacy):
        os.replace(legacy, os.path.join(a.outdir, "pdf_comparison_superseded.csv"))
        print(f"  reviews moved from {LEGACY_CSV} into {WORKBOOK}; the old CSV was renamed "
              f"pdf_comparison_superseded.csv", file=sys.stderr)

    n = {s: sum(r["status"] == s for r in rows) for s in STATUS_ORDER}
    new_gap = sum(r["status"] == "NEEDS_ARCHIVING" and r["new_this_run"] == "yes" for r in rows)
    todo = sum(r["status"] == "NEEDS_ARCHIVING" and r["appraisal"] not in APPRAISAL_VALUES for r in rows)
    print(f"\nLive PDFs: {len(pdfs)}  |  Archived PDFs: {len(arch)}", file=sys.stderr)
    print(f"NEEDS_ARCHIVING: {n['NEEDS_ARCHIVING']} ({new_gap} new this run, {todo} awaiting appraisal)  "
          f"ARCHIVED: {n['ARCHIVED']}  ARCHIVED_NOT_LIVE: {n['ARCHIVED_NOT_LIVE']}", file=sys.stderr)
    print(f"Review file: {wb_path}", file=sys.stderr)


if __name__ == "__main__":
    main()
